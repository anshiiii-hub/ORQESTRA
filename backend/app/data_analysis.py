"""Bounded, deterministic CSV/XLSX analysis; spreadsheet expressions are never evaluated."""

import csv
import hashlib
import html
import io
import math
import platform
import zipfile
from dataclasses import dataclass
from importlib.metadata import version
from pathlib import PurePath
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_XLSX_BYTES = 25 * 1024 * 1024
MAX_ZIP_MEMBERS = 2_000
MAX_ROWS = 5_000
MAX_COLUMNS = 100
MAX_CELL_CHARS = 16_384
MAX_CHART_COLUMNS = 12
ANALYSIS_VERSION = "1.0.0"


class AnalysisError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class AnalysisResult:
    filename: str
    file_format: str
    source_size_bytes: int
    content_sha256: str
    row_count: int
    column_count: int
    metadata: dict[str, Any]
    profile: dict[str, Any]
    chart_artifacts: list[dict[str, str]]
    normalized_csv: str


def safe_filename(filename: str) -> str:
    if (
        not filename
        or len(filename) > 255
        or any(ord(char) < 32 or ord(char) == 127 for char in filename)
        or "/" in filename
        or "\\" in filename
    ):
        raise AnalysisError("filename_invalid", "Choose a file with a simple name")
    safe = PurePath(filename).name
    suffix = PurePath(safe).suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        raise AnalysisError("file_type_unsupported", "Choose a CSV or XLSX file")
    return safe


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_MEMBERS:
                raise AnalysisError(
                    "xlsx_archive_too_large", "The workbook contains too many parts"
                )
            total_size = 0
            for info in infos:
                normalized = info.filename.replace("\\", "/")
                if normalized.startswith("/") or ".." in normalized.split("/"):
                    raise AnalysisError("xlsx_archive_invalid", "The workbook archive is invalid")
                if info.flag_bits & 0x1:
                    raise AnalysisError(
                        "xlsx_archive_invalid", "Encrypted workbooks are not supported"
                    )
                if normalized.lower().endswith("vbaproject.bin"):
                    raise AnalysisError(
                        "macros_unsupported", "Macro-enabled workbooks are not supported"
                    )
                total_size += info.file_size
                if total_size > MAX_EXPANDED_XLSX_BYTES:
                    raise AnalysisError(
                        "xlsx_expanded_too_large", "The workbook expands beyond the size limit"
                    )
                if info.file_size > 1024 * 1024 and info.compress_size == 0:
                    raise AnalysisError("xlsx_archive_invalid", "The workbook archive is invalid")
                if info.compress_size and info.file_size / info.compress_size > 500:
                    raise AnalysisError(
                        "xlsx_expanded_too_large", "The workbook compression ratio is too high"
                    )
            if (
                "[Content_Types].xml" not in archive.namelist()
                or "xl/workbook.xml" not in archive.namelist()
            ):
                raise AnalysisError("xlsx_invalid", "The file is not a valid XLSX workbook")
    except AnalysisError:
        raise
    except (OSError, zipfile.BadZipFile, ValueError):
        raise AnalysisError("xlsx_invalid", "The file is not a valid XLSX workbook") from None


def _csv_frame(content: bytes) -> tuple[pd.DataFrame, dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig", errors="strict")
        sample = text[:32_768]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        header = next(reader, None)
        if not header or len(header) > MAX_COLUMNS:
            raise AnalysisError(
                "csv_schema_invalid", "The CSV needs a header and no more than 100 columns"
            )
        for index, row in enumerate(reader, start=1):
            if index > MAX_ROWS:
                raise AnalysisError("dataset_row_limit", "The dataset exceeds the 5,000 row limit")
            if len(row) > MAX_COLUMNS:
                raise AnalysisError(
                    "dataset_column_limit", "The dataset exceeds the 100 column limit"
                )
            if any(len(cell) > MAX_CELL_CHARS for cell in row):
                raise AnalysisError(
                    "dataset_cell_too_large", "A cell exceeds the 16,384 character limit"
                )
        frame = pd.read_csv(
            io.BytesIO(content),
            dtype=object,
            keep_default_na=False,
            na_filter=False,
            encoding="utf-8-sig",
            sep=dialect.delimiter,
            on_bad_lines="error",
        )
    except AnalysisError:
        raise
    except (UnicodeDecodeError, csv.Error, pd.errors.ParserError, ValueError):
        raise AnalysisError(
            "csv_parse_failed", "The CSV could not be parsed as UTF-8 tabular data"
        ) from None
    return frame, {"delimiter": dialect.delimiter, "encoding": "utf-8-sig"}


def _xlsx_frame(content: bytes) -> tuple[pd.DataFrame, dict[str, Any]]:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        worksheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if not worksheets:
            raise AnalysisError("xlsx_sheet_missing", "The workbook has no visible worksheet")
        sheet = worksheets[0]
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
            raise AnalysisError(
                "dataset_dimensions_exceeded", "The worksheet exceeds the row or column limit"
            )
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or not rows[0]:
            raise AnalysisError("dataset_empty", "The worksheet has no tabular data")
        header = list(rows[0])
        if len(header) > MAX_COLUMNS:
            raise AnalysisError("dataset_column_limit", "The dataset exceeds the 100 column limit")
        if len(rows) - 1 > MAX_ROWS:
            raise AnalysisError("dataset_row_limit", "The dataset exceeds the 5,000 row limit")
        for row in rows:
            if any(isinstance(cell, str) and len(cell) > MAX_CELL_CHARS for cell in row):
                raise AnalysisError(
                    "dataset_cell_too_large", "A cell exceeds the 16,384 character limit"
                )
        frame = pd.DataFrame(rows[1:], columns=header)
        metadata = {"worksheet": sheet.title, "visible_worksheet_count": len(worksheets)}
        workbook.close()
        return frame, metadata
    except AnalysisError:
        raise
    except Exception:
        raise AnalysisError(
            "xlsx_parse_failed", "The workbook could not be safely parsed"
        ) from None


def _normalized_columns(frame: pd.DataFrame) -> pd.DataFrame:
    names: list[str] = []
    used: set[str] = set()
    for index, value in enumerate(frame.columns):
        base = str(value).strip() if value is not None else ""
        if not base:
            base = f"column_{index + 1}"
        name = base[:200]
        candidate = name
        suffix = 2
        while candidate.casefold() in used:
            ending = f"_{suffix}"
            candidate = f"{name[: 200 - len(ending)]}{ending}"
            suffix += 1
        names.append(candidate)
        used.add(candidate.casefold())
    frame.columns = names
    return frame


def _clean_value(value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _column_profile(frame: pd.DataFrame) -> tuple[list[dict[str, Any]], list[tuple[str, float]]]:
    profiles = []
    means: list[tuple[str, float]] = []
    for column in frame.columns:
        series = frame[column]
        cleaned = series.map(_clean_value)
        present = cleaned.map(lambda value: value is not None and str(value).strip() != "")
        non_missing = cleaned[present]
        numeric = pd.to_numeric(non_missing, errors="coerce")
        numeric_count = int(numeric.notna().sum())
        finite_count = int(
            numeric.map(
                lambda value: math.isfinite(float(value)) if pd.notna(value) else False
            ).sum()
        )
        is_numeric = (
            bool(len(non_missing))
            and numeric_count == len(non_missing)
            and finite_count == numeric_count
        )
        stats = numeric.astype(float) if is_numeric else None
        if stats is not None:
            with np.errstate(over="ignore", invalid="ignore"):
                mean = float(stats.mean())
                total = float(stats.sum())
            if not math.isfinite(mean) or not math.isfinite(total):
                is_numeric = False
                stats = None
        distinct = int(non_missing.nunique(dropna=True))
        item: dict[str, Any] = {
            "name": str(column),
            "inferred_type": "numeric" if is_numeric else "text",
            "missing_count": int(len(cleaned) - int(present.sum())),
            "distinct_count": distinct,
        }
        if len(non_missing):
            item["example"] = str(non_missing.iloc[0])[:160]
        if is_numeric:
            assert stats is not None
            item["statistics"] = {
                "count": numeric_count,
                "mean": round(mean, 8),
                "min": round(float(stats.min()), 8),
                "max": round(float(stats.max()), 8),
                "sum": round(total, 8),
            }
            if math.isfinite(mean):
                means.append((str(column), mean))
        profiles.append(item)
    return profiles, means[:MAX_CHART_COLUMNS]


def _chart_svg(means: list[tuple[str, float]]) -> str:
    width, height, left, right, top, row_height = 760, 100, 210, 20, 18, 34
    height = top * 2 + max(len(means), 1) * row_height
    plot_width = width - left - right
    min_value = min(0.0, min((value for _, value in means), default=0.0))
    max_value = max(0.0, max((value for _, value in means), default=1.0))
    span = max_value - min_value or 1.0
    baseline = left + (-min_value / span) * plot_width
    bars = []
    for index, (label, value) in enumerate(means):
        y = top + index * row_height
        bar_width = round(abs(value) / span * (plot_width - 70), 2)
        x = round(baseline if value >= 0 else baseline - bar_width, 2)
        text = html.escape(label, quote=True)
        bars.append(
            f'<text x="{left - 10}" y="{y + 17}" text-anchor="end">{text}</text>'
            f'<rect x="{x}" y="{y}" width="{bar_width}" height="22" rx="4" fill="#4ea2b7" />'
            f'<text x="{x + bar_width + 8}" y="{y + 17}">{value:.6g}</text>'
        )
    body = "".join(bars) or '<text x="20" y="40">No numeric columns available</text>'
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Mean by numeric column" '
        f'viewBox="0 0 {width} {height}" width="{width}" height="{height}">'
        f"<style>text{{font:12px sans-serif;fill:#d7e2ec}}</style>"
        f'<text x="20" y="16">Mean by numeric column</text>{body}</svg>'
    )


def _safe_export_cell(value: Any) -> Any:
    cleaned = _clean_value(value)
    if isinstance(cleaned, str):
        candidate = cleaned.lstrip(" \t\r\n")
        if candidate.startswith("-"):
            try:
                if math.isfinite(float(candidate)):
                    return cleaned
            except ValueError:
                pass
        if candidate.startswith(("=", "+", "-", "@")):
            return "'" + cleaned
    return cleaned


def analyze_dataset(filename: str, content: bytes) -> AnalysisResult:
    safe_name = safe_filename(filename)
    if not content:
        raise AnalysisError("file_empty", "The selected file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise AnalysisError("file_too_large", "The upload exceeds the 5 MiB limit")
    suffix = PurePath(safe_name).suffix.lower()
    if suffix == ".csv":
        frame, parser_metadata = _csv_frame(content)
        file_format = "csv"
    else:
        frame, parser_metadata = _xlsx_frame(content)
        file_format = "xlsx"
    frame = _normalized_columns(frame)
    if frame.empty:
        raise AnalysisError("dataset_empty", "The dataset has no tabular data")
    if len(frame.columns) > MAX_COLUMNS or len(frame) > MAX_ROWS:
        raise AnalysisError(
            "dataset_dimensions_exceeded", "The dataset exceeds its row or column limit"
        )
    if not len(frame.columns):
        raise AnalysisError("dataset_empty", "The dataset has no columns")
    profiles, means = _column_profile(frame)
    cleaned_rows = [
        [_clean_value(value) for value in row] for row in frame.itertuples(index=False, name=None)
    ]
    safe_frame = pd.DataFrame(
        [
            [_safe_export_cell(value) for value in row]
            for row in frame.itertuples(index=False, name=None)
        ],
        columns=frame.columns,
    )
    normalized_csv = safe_frame.to_csv(index=False, lineterminator="\n")
    metadata = {
        "analysis_version": ANALYSIS_VERSION,
        "pandas_version": pd.__version__,
        "openpyxl_version": version("openpyxl"),
        "python_version": platform.python_version(),
        "source_size_bytes": len(content),
        "source_sha256": hashlib.sha256(content).hexdigest(),
        "file_format": file_format,
        "parser": parser_metadata,
        "formula_cells_treated_as_data": sum(
            isinstance(value, str) and value.startswith("=")
            for row in cleaned_rows
            for value in row
        ),
    }
    chart_svg = _chart_svg(means)
    return AnalysisResult(
        filename=safe_name,
        file_format=file_format,
        source_size_bytes=len(content),
        content_sha256=hashlib.sha256(content).hexdigest(),
        row_count=len(frame),
        column_count=len(frame.columns),
        metadata=metadata,
        profile={"rows": len(frame), "columns": profiles},
        chart_artifacts=[
            {"id": "numeric-means", "title": "Mean by numeric column", "svg": chart_svg}
        ],
        normalized_csv=normalized_csv,
    )


def export_csv(csv_text: str) -> str:
    # Kept as a separate typed adapter to make the stored, safe output directly reusable.
    return csv_text


def validate_upload_filename(value: str | None) -> str:
    if value is None:
        raise AnalysisError("filename_missing", "Provide a file name")
    return safe_filename(value)
